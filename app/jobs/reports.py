"""Daily summary, weekly Excel report, stale lead reminder. Pure functions over sheet records."""

import io
from collections import Counter
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill

OPEN_STATUSES = {"Новая", "New", ""}


def _when(record: dict) -> datetime | None:
    value = record.get("Дата")
    if isinstance(value, datetime):
        return value
    try:
        return datetime.strptime(str(value), "%Y-%m-%d %H:%M")
    except ValueError:
        return None


def leads_on(records: list[dict], day: date) -> list[dict]:
    return [r for r in records if (w := _when(r)) and w.date() == day]


def daily_summary(records: list[dict], day: date) -> str:
    today = leads_on(records, day)
    by_source = Counter(str(r.get("Источник") or "—") for r in today)
    open_total = sum(1 for r in records if str(r.get("Статус") or "") in OPEN_STATUSES)
    lines = [f"📊 <b>Summary for {day:%d.%m.%Y}</b>", f"Leads today: {len(today)}"]
    lines += [f"  • {src}: {n}" for src, n in by_source.most_common()]
    lines.append(f"Not processed (new), all time: {open_total}")
    return "\n".join(lines)


def stale_leads(records: list[dict], now: datetime, hours: int) -> list[dict]:
    limit = now - timedelta(hours=hours)
    return [
        r
        for r in records
        if str(r.get("Статус") or "") in OPEN_STATUSES and (w := _when(r)) and w <= limit
    ]


def stale_message(stale: list[dict], hours: int) -> str:
    lines = [f"⏰ <b>{len(stale)} lead(s) without reply for more than {hours} h</b>"]
    for r in stale[:15]:
        lines.append(f"• {r.get('Дата')} {r.get('Имя') or ''} {r.get('Телефон') or ''}".rstrip())
    if len(stale) > 15:
        lines.append(f"…and {len(stale) - 15} more")
    return "\n".join(lines)


def weekly_report(records: list[dict], week_end: date) -> tuple[bytes, dict]:
    """Excel for the 7 days ending on `week_end`: per-day chart, sources, statuses, all leads."""
    days = [week_end - timedelta(days=i) for i in range(6, -1, -1)]
    week = [r for r in records if (w := _when(r)) and days[0] <= w.date() <= week_end]
    per_day = Counter(_when(r).date() for r in week)
    by_source = Counter(str(r.get("Источник") or "—") for r in week)
    by_status = Counter(str(r.get("Статус") or "Новая") for r in week)

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F4E78")
    ws.append([f"Leads report {days[0]:%d.%m.%Y} – {week_end:%d.%m.%Y}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Day", "Leads"])
    for d in days:
        ws.append([d.strftime("%a %d.%m"), per_day.get(d, 0)])
    ws.append(["Total", len(week)])
    for cell in ws[3]:
        cell.font, cell.fill = head, fill

    chart = BarChart()
    chart.title, chart.y_axis.title = "Leads per day", "Leads"
    chart.legend = None
    chart.add_data(Reference(ws, min_col=2, min_row=3, max_row=10), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=4, max_row=10))
    chart.height, chart.width = 7, 14
    ws.add_chart(chart, "D3")

    start = 14
    for title, counter in (("Source", by_source), ("Status", by_status)):
        ws.cell(start, 1, title).font = head
        ws.cell(start, 1).fill = fill
        ws.cell(start, 2, "Leads").font = head
        ws.cell(start, 2).fill = fill
        for i, (k, v) in enumerate(counter.most_common(), start=1):
            ws.cell(start + i, 1, k)
            ws.cell(start + i, 2, v)
        start += len(counter) + 3
    ws.column_dimensions["A"].width = 22

    data = wb.create_sheet("Leads")
    cols = list(records[0].keys()) if records else ["Дата"]
    data.append(cols)
    for r in week:
        data.append([str(r.get(c) or "") for c in cols])
    for cell in data[1]:
        cell.font, cell.fill = head, fill
    data.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    stats = {"total": len(week), "per_day": {str(d): per_day.get(d, 0) for d in days}}
    return buf.getvalue(), stats
