"""Demo-mode storage: data/leads.xlsx with the same columns as the Google Sheet."""

import threading
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from app.models import COLUMNS, HEADERS, STATUS_LABELS, Lead

_lock = threading.Lock()


class LocalXlsx:
    def __init__(self, path: Path, lang: str = "ru"):
        self.path = path
        self.lang = lang

    def _create(self) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "Leads"
        ws.append(HEADERS[self.lang])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        ws.freeze_panes = "A2"
        dv = DataValidation(
            type="list", formula1='"' + ",".join(STATUS_LABELS[self.lang]) + '"', allow_blank=True
        )
        ws.add_data_validation(dv)
        dv.add("F2:F10000")
        for col, width in zip("ABCDEFGHI", (17, 20, 15, 16, 40, 12, 16, 24, 30), strict=True):
            ws.column_dimensions[col].width = width
        wb.save(self.path)

    def append(self, lead: Lead) -> str:
        with _lock:
            if not self.path.exists():
                self._create()
            wb = load_workbook(self.path)
            ws = wb.active
            ws.append(lead.row(self.lang))
            row = ws.max_row
            wb.save(self.path)
        return f"{self.path.name}, row {row}"

    def records(self) -> list[dict]:
        if not self.path.exists():
            return []
        ws = load_workbook(self.path, read_only=True).active
        rows = ws.iter_rows(min_row=2, values_only=True)
        return [dict(zip(COLUMNS, r, strict=False)) for r in rows if any(r)]
